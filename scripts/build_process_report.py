"""Append verified process evidence to the supplied result report.

Requires python-docx. No model is called. Render the resulting DOCX for visual QA.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPO_URL = 'https://github.com/CHANSingYeungSunny/Design-Molecule-Class'
FONT = 'Noto Sans CJK SC'


def font(run, size=None, bold=None):
    run.font.name = 'Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold


def paragraph(doc, text, size=10.5, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    font(p.add_run(text), size, bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        font(run, 14 if level == 1 else 11, True)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def code(doc, text, width=88):
    # Long source lines are wrapped for print only; the full file stays unchanged.
    for line in text.strip('\n').splitlines():
        chunks = textwrap.wrap(line, width, subsequent_indent='    ',
                               replace_whitespace=False, drop_whitespace=False,
                               break_long_words=False, break_on_hyphens=False) or ['']
        for chunk in chunks:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(chunk)
            run.font.name = 'DejaVu Sans Mono'
            run.font.size = Pt(9)
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for cell, value in zip(t.rows[0].cells, headers): cell.text = value
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row): cell.text = str(value)
    for row_num, row in enumerate(t.rows):
        # Permit automatic height, forbid a row being cut between pages.
        props = row._tr.get_or_add_trPr(); no_split = OxmlElement('w:cantSplit'); props.append(no_split)
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = 1
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single')
                edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + side); e.set(qn('w:w'), '90'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcpr.append(margins)
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'DCE6EC' if row_num == 0 else ('F7F9FA' if row_num % 2 else 'FFFFFF'));tcpr.append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2);p.paragraph_format.line_spacing = 1.1
                for r in p.runs: font(r, 9.5, row_num == 0)
    return t


def hyperlink(doc, label, url):
    p = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    run = OxmlElement('w:r');props = OxmlElement('w:rPr')
    c = OxmlElement('w:color');c.set(qn('w:val'), '0563C1');props.append(c)
    f = OxmlElement('w:rFonts');f.set(qn('w:ascii'), 'Calibri');f.set(qn('w:eastAsia'), FONT);props.append(f)
    sz = OxmlElement('w:sz');sz.set(qn('w:val'), '21');props.append(sz)
    u = OxmlElement('w:u');u.set(qn('w:val'), 'single');props.append(u)
    run.append(props);text = OxmlElement('w:t');text.text = label;run.append(text);link.append(run);p._p.append(link)
    return p


def new_page(doc, title):
    p = heading(doc, title)
    p.paragraph_format.page_break_before = True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT / 'docs/source_report.docx')
    parser.add_argument('--out', type=Path, default=ROOT / 'docs/分子表示比较研究简明汇报_过程与复现版.docx')
    args = parser.parse_args()
    doc = Document(args.source)
    # Keep the original result sections and all numerical cells intact.
    for p in doc.paragraphs:
        for run in p.runs: font(run)
        # Preserve the original text, while removing its forced split that
        # left only the third finding on a largely empty intervening page.
        p.paragraph_format.page_break_before = False
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(3)
        if p.style.name.startswith('Heading'):
            p.paragraph_format.space_before = Pt(6)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(1)
                    for run in p.runs: font(run)
    for name in ['Title', 'Heading 1', 'Heading 2', 'Normal']:
        style = doc.styles[name];style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = 'Calibri';style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    appendix_note = paragraph(doc,
        '本次补充：后续各页展示原论文方法、实际运行的核心代码与五项模型prompt、逐题评分证据，以及不用API Key的复算命令。已完成本地复算；GitHub发布和从远端clone验收尚待本机执行。')
    # The note remains at the end of the original report, ahead of the appendix.
    templates = json.loads((ROOT / 'prompts/benchmark_templates.json').read_text(encoding='utf-8'))

    heading(doc, '七 原论文如何进行比较')
    paragraph(doc, '论文研究的是分子的写法怎样影响LLM表现。作者先固定任务与参考分子，把同一分子转换成不同文本表示，给模型相应输入和统一任务提示，再用独立的标准答案评价。模型的理解能力由回答成绩体现。')
    paragraph(doc, '正文第3至4节介绍任务与实验；PDF第15页的Appendix C和D给出数据和任务定义；第16至17页的Appendix E公开实际prompt。本文依据这份25页PDF，不声称已经取得或运行作者的完整源码。')
    table(doc, ['比较项目', '原论文', '本项目'], [
        ['输入表示', '9种文本表示', '8种输入形式，包含Our Code和图片'],
        ['任务与模型', '8项任务，16个LLM', '5项任务，一个模型deepseek-v4.1-flash'],
        ['数据划分', 'ChEBI-20官方test；部分任务ZINC250K', 'ChEBI-20再发布validation/test合并后由本项目重划'],
        ['结构门槛', 'MW ≥ 300 Da且至少2环', '沿用门槛，并要求Our Code可解析与一致往返'],
        ['元素计数', '每次询问一种元素', '一次询问C/N/O/S，四种全部正确才算对'],
        ['官能团', '逐次询问5类中的一类', '一次问13类，固定10类参与主Macro-F1'],
        ['任务输出', '逐步推理后给boxed答案', '要求一个简短JSON对象'],
        ['性质任务', 'LogP、TPSA、HBD、HBA', '本轮HBD；新增QED估计'],
    ], [1.12, 2.40, 3.02])
    paragraph(doc, '因此，我们复用了“控制模型和题目、改变输入形式”的比较思路，并修改了任务、问法和实验范围。这是扩展实验。它不能直接继承原论文的全部结论。')
    paragraph(doc, '原文正文对LogP/TPSA使用Spearman相关；Appendix D另有exact-match表述，二者措辞不完全一致。本项目QED明确采用MAE，不能写成论文原指标。')

    new_page(doc, '八 核心代码怎样把分子变成模型输入')
    paragraph(doc, '整条流程是：原始分子 → 筛选并冻结题单 → 生成表示 → 组装prompt → 模型回答 → 对照RDKit参考答案 → 算分。模型只接收表示和问题；参考答案留在程序一侧。')
    paragraph(doc, '基础类在 research/01_implementation/molecule_core.py。Atom保存一个原子的记录，Bond保存两个原子之间的键，Molecule把这些记录组成一个分子。以下是原代码的字段摘录，省略校验方法。')
    code(doc, '''@dataclass(frozen=True)
class Atom:
    id: str
    element: str
    formal_charge: int = 0
    isotope: int = 0
    aromatic: bool = False
    explicit_hs: int | None = None
    atom_map: int | None = None

@dataclass(frozen=True)
class Bond:
    begin: str
    end: str
    bond_type: str = 'SINGLE' ''')
    paragraph(doc, 'id是分子内部标签，element是元素，formal_charge是电荷，aromatic表示是否芳香。explicit_hs=None沿用隐式氢规则；整数表示确定的显式氢数。未在构造文本写出的字段采用默认值，不能把“省略字段”理解成“没有氢”。')
    paragraph(doc, '真正的表示转换入口在bmb/reps.py，最终执行基础类的to_code。下面是该函数的实际源码，打印换行仅为排版。')
    code(doc, '''def our_code_text(smiles: str) -> str:
    """The professor prototype's constructor text for this molecule."""
    try:
        return _core_molecule(smiles).to_code()
    except Exception as exc:
        raise RepresentationError('Our Code cannot express %r: %s' % (smiles, exc)) from exc''')
    paragraph(doc, '例如乙醇CCO产生三条Atom和两条Bond。模型阅读这些构造代码，不执行Python来得到答案。这个教学分子太小，不满足正式复杂度门槛。')
    code(doc, '''Molecule(
    atoms=(Atom(id='a0', element='C'),
           Atom(id='a1', element='C'),
           Atom(id='a2', element='O')),
    bonds=(Bond(begin='a0', end='a1'),
           Bond(begin='a1', end='a2')),
)''')
    paragraph(doc, '上面乙醇构造为了阅读合并了行；完整生成格式与全部方法见源码。教授原型的三份核心文件字节一致，SHA-256以55d314aa开头；本次未改基础类。')

    new_page(doc, '九 模型实际收到的提示词')
    paragraph(doc, '以下来自实际bmb/tasks.py及冻结历史输入，不是新编的优化提示。完整模板、八种表示说明和五个真实例子已导出到prompts目录。CodeBuddy执行指令是开发工具的任务单，与这些分子题目不同。')
    heading(doc, '所有任务共用的system prompt', 2)
    code(doc, templates['system'])
    paragraph(doc, '意思是：只能根据给定结构回答，不运行代码或查询工具，只输出一个简短JSON。模型的正确与否由外部程序判断。')
    heading(doc, 'Our Code输入说明', 2)
    code(doc, templates['representation_notes']['our_code'])
    paragraph(doc, 'user消息的结构是Representation说明、Molecule分子内容、任务instruction；Identity改为A和B两份内容。省略分子内容的模板不能冒称一条完整历史请求；完整请求在historical_our_code_examples.json。')
    heading(doc, '同分子判断的实际instruction', 2)
    code(doc, templates['templates']['t1_identity']['task_instruction'])
    paragraph(doc, '模型应返回same_molecule的true或false；依据原子、键、键级和电荷判断，不比较字符串是否一样。历史输入排序规则仍有差异，所以该项不能作确认性优势结论。')
    heading(doc, '官能团识别的实际instruction', 2)
    code(doc, templates['templates']['t2_functional_groups']['task_instruction'])
    paragraph(doc, '实际问13类；主Macro-F1只对冻结的10类平均。aldehyde、nitrile、nitro不进入主分数。原示例里的双大括号是已发现的格式缺陷，这里保留实际历史问法。')

    new_page(doc, '十 其余任务怎样问与怎样判分')
    heading(doc, '元素计数的实际instruction', 2)
    code(doc, templates['templates']['t3_atom_counts']['task_instruction'])
    paragraph(doc, 'C、N、O、S四个数全部与参考相同，该题才算正确。原示例同样含双大括号；未来修正需另立prompt版本，不能改写旧运行记录。')
    heading(doc, 'QED估计的实际instruction', 2)
    code(doc, templates['templates']['t4_qed']['task_instruction'])
    paragraph(doc, '模型估计0到1的QED；程序计算回答与RDKit值的绝对差，再对题目平均得到MAE，越小越好。该值不是疗效或安全性。')
    heading(doc, '氢键供体计数的实际instruction', 2)
    code(doc, templates['templates']['t5_hydrogen_bond_donors']['task_instruction'])
    paragraph(doc, 'HBD问的是供体位点数量，不是分子内所有氢的数量。历史prompt用带氢的N/O作简化说明，oracle则用RDKit CalcNumHBD；特殊原子的定义边界应在新版本检查。')
    heading(doc, '一个真实的历史失分例子', 2)
    paragraph(doc, '题目t3-our_code-71464601取自main-02原始日志；它是元素计数任务。完整Our Code输入、模型回答与另列参考答案见prompts/historical_our_code_examples.json。')
    table(doc, ['元素', '模型回答', '参考答案', '是否一致'], [
        ['C', '22', '21', '否'], ['N', '1', '1', '是'], ['O', '4', '4', '是'], ['S', '2', '2', '是'],
    ], [1.25, 1.65, 1.65, 1.99])
    paragraph(doc, '模型多算了一个碳。因此四元素exact-match规则下，该题得0；不是给3/4分。这个例子证明如何产生失分，但不能单凭它确定为什么模型多算了碳。')

    new_page(doc, '十一 参考答案 请求与评分的核心代码')
    paragraph(doc, 'bmb/labels.py产生参考答案。它使用独立RDKit，不使用我们的基础类给自己的表示打分。以下是实际源码片段。')
    code(doc, '''def atom_count(smiles: str, element: str) -> int:
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError('RDKit cannot parse %r' % smiles)
    return sum(1 for atom in mol.GetAtoms() if atom.GetSymbol() == element)

def qed_reference(smiles: str) -> float:
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError('RDKit cannot parse %r' % smiles)
    return float(QED.qed(mol))''')
    paragraph(doc, '官能团用mol.HasSubstructMatch匹配冻结SMARTS模式；HBD用rdMolDescriptors.CalcNumHBD。SMARTS是指定局部原子和键组合的查询规则，定义在config/functional_groups.frozen.json。')
    paragraph(doc, 'bmb/client.py的_payload把prompt放入真实请求，以下是实际请求体构造片段。extra_body中的enable_thinking设置以及response_format也由该函数加入。')
    code(doc, '''body = {
    'model': self.model_id,
    'messages': messages,
    'temperature': self.temperature,
    'top_p': self.top_p,
    'max_tokens': self.max_tokens,
    'stream': False,
}''')
    paragraph(doc, '历史记录为deepseek-v4.1-flash、temperature=0、top_p=1、max_tokens=4096、zero-shot，请求enable_thinking=false。run manifest实际记录Token Plan通道；本次只离线复算，没有新调用模型。服务端未提供固定权重版本。')
    paragraph(doc, 'bmb/scoring.py的score_item直接比较回答和参考。以下为函数中的真实分支；前面的status校验与其它分支在完整源码。')
    code(doc, '''if task == 't3_atom_counts':
    return bool(parsed == label)
if task == 't4_qed':
    return abs(parsed['qed'] - label['qed'])
if task == 't5_hydrogen_bond_donors':
    return bool(parsed['hbd'] == label['hbd'])''')
    paragraph(doc, '子结构分数不是逐题exact-match：程序分别累计每类的正确识别、误报和漏报，计算F1，再对固定类别等权平均。共同题表由build_common_table.py先按每项任务取八种表示的交集，再重算指标。')

    new_page(doc, '十二 怎样复算并从GitHub审查')
    paragraph(doc, '最短复算路径只需要Python 3.10或更高版本。先解压本次发布包，进入包含README.md的目录，再执行：')
    code(doc, '''python research/02_tests/test_baseline_core.py
python scripts/replay_table.py''')
    paragraph(doc, '第一条验证原子、键和分子基础类；第二条重新解析保存的模型JSON，用原评分器重算主表。输出在reproduced/：summary.csv是表格，per_item.jsonl是逐题评分证据，verification.json是校验结果。参考标签采用当时保存的RDKit值；这一步不是重新调用模型。')
    paragraph(doc, '本次已核验：核心55/55、benchmark自检99/99、客户端离线检查14/14；40份冻结题单、main的3000项及extension的4996项输入hash通过；八行40格共同题结果与原Word一致。共同题分母仍为200/199/199/199/199。')
    paragraph(doc, '完整环境、标签与输入重建需要RDKit等依赖；历史环境版本和安装命令在README。新模型调用需自备服务、key和预算，并另用run ID。新调用可能得到不同回答，不能承诺重新推理逐题一致。')
    heading(doc, 'GitHub目标仓库', 2)
    hyperlink(doc, REPO_URL, REPO_URL)
    paragraph(doc, '仓库地址由我提供。本次已整理本地发布包；当前环境不能读取或推送该远端，尚未确认线上包含这些文件、未取得发布commit或完成fresh clone验收。以下超链接是预定main路径，发布后应改为实际commit固定链接并核验。')
    for label, path in [
        ('README 运行步骤与结果', 'README.md'),
        ('基础类 Atom Bond Molecule', 'research/01_implementation/molecule_core.py'),
        ('实际prompt builder', 'research/07_benchmark/bmb/tasks.py'),
        ('导出的完整prompt模板', 'prompts/benchmark_templates.json'),
        ('独立RDKit参考答案', 'research/07_benchmark/bmb/labels.py'),
        ('评分规则', 'research/07_benchmark/bmb/scoring.py'),
        ('无API Key历史复算入口', 'scripts/replay_table.py'),
    ]:
        hyperlink(doc, label, REPO_URL + '/blob/main/' + path)
    paragraph(doc, 'CodeBuddy下一步应安全合并本地修改、推送这个仓库，从新目录clone实际commit并按README验收，然后把真实commit、固定链接和发布状态更新到报告。只有完成这些步骤，才能写成“其他人已可从GitHub复现”。')
    paragraph(doc, '我向教授说明时应能解释：谁算标准答案、模型收到什么、是否执行代码、一个回答如何算分、哪些改动来自论文、哪些是项目选择，以及现有比较限制。提交代码有助于审核，但理解这些过程仍需要我自己完成。')

    args.out.parent.mkdir(parents=True, exist_ok=True)
    # The supplied file used a CJK document grid. Snap-to-grid inflated the
    # actual line heights of code and made short trailing paragraphs spill.
    for section in doc.sections:
        for grid in list(section._sectPr.findall(qn('w:docGrid'))):
            section._sectPr.remove(grid)
    for p in doc.element.iter(qn('w:p')):
        props = p.find(qn('w:pPr'))
        if props is None:
            props = OxmlElement('w:pPr'); p.insert(0, props)
        for old in list(props.findall(qn('w:snapToGrid'))): props.remove(old)
        snap = OxmlElement('w:snapToGrid');snap.set(qn('w:val'), '0');props.append(snap)
    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    main()
